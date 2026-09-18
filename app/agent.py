"""Experimental local tool-calling agent (off by default; see ENABLE_AGENT_MODE).

Runs entirely against the local RAG index and local model — no network calls.
This is a prototype for the "agents on government terminals" roadmap item:
a small ReAct-style loop where the model can search the document library,
read a specific file, and draft a Word report, grounding its final answer in
what the tools actually returned rather than in memory.

Requires a chat model that supports OpenAI-style tool calling (most current
Ollama chat models do — qwen2.5, llama3.1+, mistral-nemo, etc). If the
configured model doesn't support tools, the underlying API call will raise;
callers should surface that error to the user rather than silently falling
back, so it's obvious the model needs to be swapped.
"""
from __future__ import annotations

import json
import re
import time
import uuid
from collections.abc import Callable
from typing import Any

from app.config import settings
from app.llm import chat_with_tools
from app.rag import get_rag

AGENT_SYSTEM_PROMPT = """You are a local, offline agent for Project SPK running entirely on this laptop \
— no internet, no OpenAI API, nothing leaves this machine. You help a USACE user work with the documents \
indexed in their local library.

You have tools to search the library, list what's indexed, read one document in full, and draft a Word \
(.docx) report file. Use a tool whenever the question depends on the documents rather than general \
knowledge. Ground every factual claim in what a tool actually returned — never invent document contents, \
page numbers, or file names. If a tool returns no results or an error, say so plainly instead of guessing.

When you're done gathering information, give your final answer as plain text with no further tool call. \
If you drafted a file, tell the user its exact filename and that it was saved to the agent_output folder \
on this laptop."""

TOOLS: list[dict[str, Any]] = [
    {
        "type": "function",
        "function": {
            "name": "search_documents",
            "description": (
                "Semantic search over the locally indexed document library. Returns the most "
                "relevant snippets with source filename and page numbers."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "What to search for."},
                    "top_k": {
                        "type": "integer",
                        "description": "Max results to return (default 8, max 20).",
                    },
                },
                "required": ["query"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "list_documents",
            "description": (
                "List every document currently indexed in the local library, with filename, "
                "document number, title, and category."
            ),
            "parameters": {"type": "object", "properties": {}},
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_document",
            "description": (
                "Read the indexed text of one specific document, in page order — use this to "
                "answer questions about a whole file (e.g. summarize it). Find the exact filename "
                "first with list_documents or search_documents."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "source": {
                        "type": "string",
                        "description": "Exact filename as shown by list_documents/search_documents.",
                    },
                    "max_chunks": {
                        "type": "integer",
                        "description": "Max sections to read (default 20, max 60).",
                    },
                },
                "required": ["source"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "draft_docx_report",
            "description": (
                "Write a Word (.docx) document to the local agent_output folder on this laptop. "
                "Use for memos, summaries, or reports the user asks you to produce as a file."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {
                        "type": "string",
                        "description": "Document title (used for the filename and heading).",
                    },
                    "sections": {
                        "type": "array",
                        "description": "Ordered list of sections.",
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string"},
                                "body": {
                                    "type": "string",
                                    "description": "Paragraph text; use a blank line between paragraphs.",
                                },
                            },
                            "required": ["heading", "body"],
                        },
                    },
                },
                "required": ["title", "sections"],
            },
        },
    },
]

MAX_SNIPPET_CHARS = 900
_FILENAME_SAFE = re.compile(r"[^\w.\- ]")


def _truncate(text: str, limit: int = MAX_SNIPPET_CHARS) -> str:
    text = text or ""
    return text if len(text) <= limit else text[:limit].rstrip() + "…"


def _tool_search_documents(args: dict[str, Any]) -> dict[str, Any]:
    query = str(args.get("query") or "").strip()
    if not query:
        return {"error": "query is required"}
    top_k = min(max(int(args.get("top_k") or 8), 1), 20)
    chunks = get_rag().retrieve(query, top_k=top_k)
    return {
        "results": [
            {
                "source": c.get("source"),
                "page_start": c.get("page_start"),
                "page_end": c.get("page_end"),
                "text": _truncate(c.get("text", "")),
            }
            for c in chunks[:top_k]
        ]
    }


def _tool_list_documents(_: dict[str, Any]) -> dict[str, Any]:
    docs = get_rag().list_documents()
    return {
        "documents": [
            {
                "source": d.get("source"),
                "doc_number": d.get("doc_number"),
                "title": d.get("title") or d.get("display_title"),
                "category": d.get("category"),
                "chunks": d.get("chunks"),
            }
            for d in docs[:300]
        ]
    }


def _tool_read_document(args: dict[str, Any]) -> dict[str, Any]:
    source = str(args.get("source") or "").strip()
    if not source:
        return {"error": "source is required"}
    max_chunks = min(max(int(args.get("max_chunks") or 20), 1), 60)
    chunks = get_rag().get_source_chunks(source, max_chunks=max_chunks)
    if not chunks:
        return {
            "error": (
                f"No indexed content found for source '{source}'. Check the exact filename "
                "with list_documents or search_documents."
            )
        }
    return {
        "source": source,
        "chunks": [
            {
                "page_start": c.get("page_start"),
                "page_end": c.get("page_end"),
                "text": _truncate(c.get("text", ""), 1600),
            }
            for c in chunks
        ],
    }


def _safe_output_name(title: str) -> str:
    base = _FILENAME_SAFE.sub("_", title).strip() or "agent_report"
    return base[:120]


def _tool_draft_docx_report(args: dict[str, Any]) -> dict[str, Any]:
    from docx import Document  # local import: only needed when this tool actually runs

    title = str(args.get("title") or "Agent Report").strip()
    sections = args.get("sections") or []
    if not isinstance(sections, list) or not sections:
        return {"error": "sections must be a non-empty list of {heading, body}"}

    doc = Document()
    doc.add_heading(title, level=0)
    for section in sections:
        if not isinstance(section, dict):
            continue
        heading = str(section.get("heading") or "").strip()
        body = str(section.get("body") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        for para in body.split("\n\n"):
            para = para.strip()
            if para:
                doc.add_paragraph(para)

    out_dir = settings.data_path / "_agent_output"
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{_safe_output_name(title)}_{uuid.uuid4().hex[:8]}.docx"
    out_path = out_dir / filename
    doc.save(out_path)
    return {"saved_as": filename, "path": str(out_path)}


TOOL_IMPLS: dict[str, Callable[[dict[str, Any]], dict[str, Any]]] = {
    "search_documents": _tool_search_documents,
    "list_documents": _tool_list_documents,
    "read_document": _tool_read_document,
    "draft_docx_report": _tool_draft_docx_report,
}


def _summarize_result(tool: str, result: dict[str, Any]) -> str:
    if "error" in result:
        return f"error: {result['error']}"
    if tool == "search_documents":
        return f"{len(result.get('results', []))} result(s)"
    if tool == "list_documents":
        return f"{len(result.get('documents', []))} document(s) indexed"
    if tool == "read_document":
        return f"{len(result.get('chunks', []))} section(s) from {result.get('source')}"
    if tool == "draft_docx_report":
        return f"saved {result.get('saved_as')}"
    return "done"


def _history_messages(history: list[dict[str, str]] | None, limit: int = 8) -> list[dict[str, Any]]:
    if not history:
        return []
    out: list[dict[str, Any]] = []
    for item in history[-limit:]:
        role = item.get("role")
        content = (item.get("content") or item.get("text") or "").strip()
        if role in ("user", "assistant") and content:
            out.append({"role": role, "content": content})
    return out


def run_agent(
    question: str,
    history: list[dict[str, str]] | None = None,
    *,
    on_step: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    """Run a tool-calling agent loop against the local index.

    Returns {"answer": str, "steps": [...], "sources": [...]}.
    `on_step` is called after each tool execution so callers can surface
    live progress (e.g. into a Job the UI is polling).
    """
    messages: list[dict[str, Any]] = [{"role": "system", "content": AGENT_SYSTEM_PROMPT}]
    messages.extend(_history_messages(history))
    messages.append({"role": "user", "content": question})

    steps: list[dict[str, Any]] = []
    sources_seen: set[str] = set()

    for _ in range(max(1, settings.agent_max_steps)):
        message = chat_with_tools(messages, TOOLS)
        tool_calls = list(getattr(message, "tool_calls", None) or [])
        if not tool_calls:
            answer = (message.content or "").strip()
            return {"answer": answer, "steps": steps, "sources": sorted(sources_seen)}

        # Preserve the assistant turn (with its tool calls) before the tool results,
        # so the model can see its own prior calls on the next round.
        messages.append(
            {
                "role": "assistant",
                "content": message.content or "",
                "tool_calls": [
                    {
                        "id": tc.id,
                        "type": "function",
                        "function": {"name": tc.function.name, "arguments": tc.function.arguments},
                    }
                    for tc in tool_calls
                ],
            }
        )

        for tc in tool_calls:
            name = tc.function.name
            try:
                args = json.loads(tc.function.arguments or "{}")
            except (json.JSONDecodeError, TypeError):
                args = {}

            impl = TOOL_IMPLS.get(name)
            t0 = time.time()
            if impl is None:
                result: dict[str, Any] = {"error": f"Unknown tool '{name}'"}
            else:
                try:
                    result = impl(args)
                except Exception as exc:  # noqa: BLE001 — surface to the model, don't crash the agent
                    result = {"error": str(exc)}
            elapsed_ms = int((time.time() - t0) * 1000)

            for bucket in ("results", "chunks"):
                for item in result.get(bucket) or []:
                    source = item.get("source") if isinstance(item, dict) else None
                    if source:
                        sources_seen.add(source)
            if result.get("source"):
                sources_seen.add(result["source"])

            step = {
                "tool": name,
                "args": args,
                "elapsed_ms": elapsed_ms,
                "ok": "error" not in result,
                "summary": _summarize_result(name, result),
            }
            steps.append(step)
            if on_step:
                on_step(step)

            messages.append(
                {
                    "role": "tool",
                    "tool_call_id": tc.id,
                    "content": json.dumps(result)[:8000],
                }
            )

    # Ran out of steps — ask once more without tools so the user still gets an answer
    # instead of silently looping forever.
    messages.append(
        {
            "role": "user",
            "content": (
                "You've used all available tool calls for this turn. Give your best final "
                "answer now, in plain text, based on what you've found so far."
            ),
        }
    )
    message = chat_with_tools(messages, [])
    return {
        "answer": (message.content or "").strip(),
        "steps": steps,
        "sources": sorted(sources_seen),
    }
